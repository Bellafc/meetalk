import re
import pandas as pd
from sentence_transformers import SentenceTransformer
from sentence_transformers import SentenceTransformer,util
import json
import dspy
from dspy.teleprompt import LabeledFewShot
import copy
from openai import OpenAI
from torch import Tensor
from transformers import AutoTokenizer, AutoModel
from docx import Document
import json

class Assess(dspy.Signature):
    """Which part does this content belong to? 优先label到部分目录中的章节去。Strictly label in this format and don't output anything else since your answer will be directly use, and don't put any quotes inside. Even there's only one label, repeat twice:Chapterxx,Sectionxx,和Chapterxx,Sectionxx"""
    part_chapters = dspy.InputField(desc="部分目录")
    pre_chapter = dspy.InputField(desc="上一个部分的章节")
    current_content = dspy.InputField(desc="给定的内容")
    label = dspy.OutputField(desc="Chapterxx,Sectionxx,和Chapterxx,Sectionxx")
class Predict(dspy.Module):
    def __init__(self):
        super().__init__()
        self.prog = dspy.Predict(Assess)
    def forward(self, part_chapters, pre_chapter, current_content):
        return self.prog(part_chapters=part_chapters, pre_chapter=pre_chapter, current_content=current_content)

def get_font_size(run):
    """
    获取文本字号
    """
    if run.font.size:
        return run.font.size.pt
    return None
def extract_tag_data(writing_style):
    # 检查writing_style是否为列表，如果是，获取第一个元素
    if isinstance(writing_style, list) and len(writing_style) > 0:
        writing_style = writing_style[0]

    # 确保writing_style是字典
    if isinstance(writing_style, dict):
        return writing_style.get("tag", "")

    # 如果传入的数据结构不对，返回空字符串
    return ""

def split_into_sentences(text):
    sentences = text.split('. ')
    return ['.'.join(sentences[i:i+2]).strip() + '.' for i in range(0, len(sentences), 2)]

def remove_duplicates(writingstyledataList):
    unique_styles = []
    seen = set()

    for item in writingstyledataList:
        # 将每个字典对象转换为字符串进行比较
        item_str = json.dumps(item, sort_keys=True)
        if item_str not in seen:
            seen.add(item_str)
            unique_styles.append(item)

    return unique_styles

def process_document(document):
    writingstyledataList=[]
    structure = {
        "writingstyledataList": [],
        "chapterList": [],
        "allocationdataList": []
    }

    current_chapter = None
    current_section = None
    current_section_list = []
    allocationdataList=[]
    last_writing_style = None

    for para in document.paragraphs:
        text = para.text.strip()
        gentext = f"""
                            analyzing this paragraph {text} and then output the writing style json as follows

                            "clout": ,
                            "input": to describe the text feature in your input file e.g. incontinuous conversation of xxx,
                            "writing_goal": write for whom for what ,
                            "writing_format": paragraph?bullet points? or what,
                            "analytical_thinking": ,
                            "authentic": ,
                            "difference": keep it empty,
                            "language": ,
                            "tag": ,
                            "emotional_tone": ,
                            "participant": ,
                            "your_role": what perspective to take, participant? observer? 

                            noted the tag should be only one to two words.directly output the json starting from a bracket and ending with a bracket. Don't output anything else 
                            """
        if not text:
            continue

        if para.runs:
            font_size = get_font_size(para.runs[0])

            if font_size == 20:
                if current_chapter:
                    structure["chapterList"].append({
                        "chapter": current_chapter,
                        "sectionList": current_section_list
                    })
                current_chapter = text
                current_section_list = []

            elif font_size == 16:
                current_section = text
                current_section_list.append({
                    "writingStyle": "default",
                    "section": current_section
                })

            else:
                sentences=split_into_sentences(text)
                for sentence in sentences:
                    structure['allocationdataList'].append({
                        "current_content": sentence,
                        "label1": f"Chapter{current_chapter},Section{current_section}",
                        "label2": f"Chapter{current_chapter},Section{current_section}"
                    })
                if last_writing_style:
                    prompt = f"Can the following writing style be applied to this paragraph?\n\nWriting Style:\n{last_writing_style}\n\nParagraph:\n{text}\n\nAnswer only 'Yes' or 'No'"
                    response = llmwriting(prompt)
                    if 'Yes' in response:
                        writing_style = last_writing_style
                    else:
                        writing_style = llmwriting(gentext)
                        last_writing_style = writing_style
                else:
                    writing_style = llmwriting(gentext)
                    last_writing_style = writing_style

                try:
                    writing_style_json = json.loads(writing_style)
                    print("writing style",writing_style_json)
                    print("writing style 0",writing_style_json[0])
                    writingstyledataList.append(writing_style_json[0])
                    writingstyledataList= remove_duplicates(writingstyledataList)
                    structure["writingstyledataList"]=writingstyledataList
                except json.JSONDecodeError:
                    print(f"Error decoding writing style: {writing_style}")

                if current_section_list:
                    current_section_list[-1]["writingStyle"] = extract_tag_data(writing_style_json)

    if current_chapter:
        structure["chapterList"].append({
            "chapter": current_chapter,
            "sectionList": current_section_list
        })

    return structure

def dataframe_to_spaced_string(df: pd.DataFrame) -> str:
    columns = df.columns.tolist()
    formatted_text = ''
    for index, row in df.iterrows():
        row_string = '\t'.join([f"{col}: {row[col]}" for col in columns])
        formatted_text += row_string + '\n'
    return formatted_text
def create_trainset(file_path):
    df = pd.read_csv(file_path, header=0)
    print(df)
    df.columns = ['current_content','label1','label2']
    result = df.to_json(orient='records', force_ascii=False)
    json_obj = json.loads(result)
    trainset = []
    for i in range(1, len(json_obj)):
        item = json_obj[i]
        current_content = item["current_content"]
        label1 = item["label1"]
        label_str1 = json.dumps(label1, ensure_ascii=False)
        label2 = item["label2"]
        label_str2 = json.dumps(label2, ensure_ascii=False)
        trainset.append(dspy.Example(
            current_content=current_content,
            label=label_str1 + ",和" + label_str2
        ).with_inputs('part_chapters', 'current_content'))
    return trainset

def average_pool(last_hidden_states: Tensor,
                 attention_mask: Tensor) -> Tensor:
    last_hidden = last_hidden_states.masked_fill(~attention_mask[..., None].bool(), 0.0)
    return last_hidden.sum(dim=1) / attention_mask.sum(dim=1)[..., None]
def generate_embeddings_from_trainset(trainset,tokenizer,model):
    alltrainembedd = []
    for i in range(len(trainset)):
        batch_dict = tokenizer(trainset[i].current_content, max_length=512, padding=True, truncation=True,
                               return_tensors='pt')
        embeddings = average_pool(model(**batch_dict).last_hidden_state, batch_dict['attention_mask'])
        alltrainembedd.append(embeddings)
    return alltrainembedd


def show_node(output,root,rootwithcontent,group,pre_chapter,trainset,alltrainembedd,k,tokenizer,model,auth,known,unknown):
    part_chapters = tree_to_string(root).replace("部分目录", "")
    print("partchapters",part_chapters)
    if pre_chapter == "ChapterUnknownChapter,SectionUnknownSection,和ChapterUnknownChapter,SectionUnknownSection":
        pre_chapter= ""
    current_content = group
    top=findtopk(k, trainset, current_content, alltrainembedd,tokenizer,model)
    topktrainset, devset = top[0],top[1]
    authentic=str(llmwriting("Can the segment allocate the segment into the given table of contents by referring the examples and table of content? Be confident." + "example:" + "".join(top[2]) + ".segment:" + current_content + ". Given table of contents:"+part_chapters+". Please answer with only Yes or No"))
    print("authentic",authentic)
    teleprompter = LabeledFewShot(k=2)
    optimized = teleprompter.compile(student=Predict(), trainset=topktrainset)
    answer = optimized(part_chapters=part_chapters, pre_chapter=pre_chapter, current_content=current_content).label
    answer = answer.replace("'", "").replace("\"", "").replace("Label", "").replace(":", "").replace(" ","")
    pattern = r"Chapter(.+),Section(.+),和Chapter(.+),Section(.+)"
    print("answerprev",answer)
    if not re.match(pattern, answer):
        answer = "ChapterUnknownChapter,SectionUnknownSection,和ChapterUnknownChapter,SectionUnknownSection"
    print("answer",answer)
    if authentic=='No':
        print("This is an unknown")
        unknown=unknown+1
    else:
        print("This is known")
        known=known+1
    answer0 = answer.split(",和")[0]
    if authentic=='No':
        if "[Not Sure]" not in answer0:
            answer0=answer0+"[Not Sure]"
    output = output + group + "(" + answer0 + ")"
    root = updatenote(root, answer0)
    rootwithcontent = updatenote(rootwithcontent, answer0)
    target_node = find_node(rootwithcontent, answer0.split(",")[1])
    if target_node:
        target_node.add_content(current_content)
    try:
        answer1 = answer.split(",和")[1]
        if authentic=='No':
            if "[Not Sure]" not in answer1:
                answer1=answer1+"[Not Sure]"
        if answer1!=answer0:
            output=output+" ("+answer1+")"
            root = updatenote(root, answer1)
            rootwithcontent = updatenote(rootwithcontent, answer1)
            target_node = find_node(rootwithcontent, answer1.split(",")[1])
            if target_node:
                target_node.add_content(current_content)
    except:
        print("no answer 1")
    output=output+"\n"
    print("root",tree_to_string(root))
    print("rootwithcontent",tree_to_string(rootwithcontent))
    return tree_to_string(root),tree_to_string(rootwithcontent), output,answer,auth,known,unknown

def tree_to_string(node, level=0):
    indent = "  " * level
    node_string = f'{indent}{node.name}'

    if node.content:
        node_string += f': "{node.content}"'

    for child in node.children:
        node_string += '\n' + tree_to_string(child, level + 1)

    return node_string


def get_children(node):
    children = []
    if node.children:
        for child in node.children:
            children.append(child)
            children.extend(get_children(child))
    return children

def find_node(root, target_name):
    if root.name == target_name:
        return root
    for child in root.children:
        result = find_node(child, target_name)
        if result:
            return result
    return None

def updatenote(root, answer):
    # 定义正则表达式模式
    pattern = r'Chapter([^"]*),Section([^"]*)'

    # 使用正则表达式进行匹配
    matches = re.search(pattern, answer)
    if matches:
        answerroot = "Chapter" + matches.group(1)
        answerchild = "Section" + matches.group(2)
    else:
        return root

    # 查找目标 Chapter 节点
    target_node = find_node(root, answerroot)

    if target_node:
        allchildren = [child.name for child in target_node.children]
        if answerchild not in allchildren:
            new_node = Node(answerchild)
            target_node.add_child(new_node)
    else:
        new_node = Node(answerroot)
        root.add_child(new_node)
        # 将新的 Section 节点添加到新的 Chapter 节点下
        new_child_node = Node(answerchild)
        new_node.add_child(new_child_node)

    return root

def chapter_api(function_name, function_instance_id, total_chapters, pre_chapter,current_content):
    url = "http://45.43.59.237:14064/api/record"
    payload = {
        "function_name": function_name,
        "function_instance_id": function_instance_id,
        "total_chapters": total_chapters,
        "pre_chapter":pre_chapter,
        "current_content":current_content

    }
    response = requests.post(url, json=payload)
    
    if response.status_code == 200:
        data = response.json()
        ai_output = data.get('ai_output')
        return ai_output
    else:
        print("Failed to call API. Status code:", response.status_code)
        return None
    
def findtopk(k,trainset,current_content,alltrainembedd,tokenizer,model):
    batch_dict = tokenizer(current_content, max_length=512, padding=True, truncation=True, return_tensors='pt')
    embeddings = average_pool(model(**batch_dict).last_hidden_state, batch_dict['attention_mask'])
    scores=[]
    topktrain=[]
    examplelist=[]
    val=[]
    for i in range(len(alltrainembedd)):
        scores.append(util.pytorch_cos_sim(embeddings, alltrainembedd[i]))
    score_index = {score: index for index, score in enumerate(scores)}
    sorted_scores = sorted(scores, reverse=True)
    print("sorted_scores",sorted_scores)

    top_k_indices = [score_index[score] for score in sorted_scores[:k] if score>0.55]
    other_indices = [score_index[score] for score in sorted_scores[k+1:] if score>0.55]
    for m in top_k_indices:
        topktrain.append(dspy.Example(current_content=trainset[m].current_content,
                                     label=trainset[m].label).with_inputs('part_chapters', "current_content")
        )
        examplelist.append("current_content"+trainset[m].current_content+".label"+trainset[m].label)
    for n in other_indices:
        val.append(dspy.Example(current_content=trainset[n].current_content,
                                     label=trainset[n].label).with_inputs('part_chapters',"current_content")
        )
    
    return topktrain,val,examplelist


def find_node(root, target_name):
    if root.name == target_name:
        return root
    for child in root.children:
        result = find_node(child, target_name)
        if result:
            return result
    return None



def segment_text(text, n):
    sentences = re.split(r'(?<=[。！？.?])', text)

    segments = []
    temp_segment = []

    for sentence in sentences:
        if sentence:
            temp_segment.append(sentence.strip())
            if (len(temp_segment) == n or
                    sum(1 for s in temp_segment if '？' in s) == n or
                    (len(temp_segment) == n and sum(1 for s in temp_segment if '？' in s) == n)):
                # Add current segment to segments list
                segments.append(' '.join(temp_segment))

                # Start the new temp_segment with the last sentence of the current temp_segment
                temp_segment = [temp_segment[-1]]

    if temp_segment:
        segments.append(' '.join(temp_segment))

    return segments


class Node:
    def __init__(self, name, children=None):
        self.name = name
        self.children = children if children is not None else []
        self.content=""

    def add_content(self, content):
        self.content += content

    def add_child(self, child_node):
        self.children.append(child_node)

    def __repr__(self):
        return f'Node("{self.name}", "{self.content}", {self.children})'
def convert_to_tree_allocation(result_list):
    root = Node("Part Chapters")
    chapter_dict = {}
    for item in result_list:
        if "no allocation" not in item['writingStyle']:
            chapter_name = "Chapter" + item["chapter"]
            section_name = "Section" + item["section"]
            if chapter_name not in chapter_dict:
                chapter_node = Node(chapter_name)
                chapter_dict[chapter_name] = chapter_node
                root.add_child(chapter_node)
            else:
                chapter_node = chapter_dict[chapter_name]
            section_node = Node(section_name)
            chapter_node.add_child(section_node)
    return root

def convert_to_tree(result_list):
    root = Node("Part Chapters")
    chapter_dict = {}
    for item in result_list:
        chapter_name = "Chapter" + item["chapter"]
        section_name = "Section" + item["section"]
        if chapter_name not in chapter_dict:
            chapter_node = Node(chapter_name)
            chapter_dict[chapter_name] = chapter_node
            root.add_child(chapter_node)
        else:
            chapter_node = chapter_dict[chapter_name]
        section_node = Node(section_name)
        chapter_node.add_child(section_node)
    return root
def difference(A,B):
    C=""
    aseg=segment_text(A,2)
    bseg=segment_text(B,2)
    for i in bseg:
        if i not in aseg:
            C+=i
    return C


def extract_sections(text):
    lines = text.split('\n')
    structure = {}
    current_chapter = ""
    current_section = ""
    for line in lines:
        line = line.strip()
        if line.startswith('**'):
            current_chapter = line
            if current_chapter not in structure:
                structure[current_chapter] = {}
            current_section = ""
        elif line.startswith('*'):
            current_section = line
            if current_chapter:
                if current_section not in structure[current_chapter]:
                    structure[current_chapter][current_section] = ""
        else:
            if current_chapter and current_section:
                structure[current_chapter][current_section] += line + " "

    # Clean up extra spaces
    for chapter in structure:
        for section in structure[chapter]:
            structure[chapter][section] = structure[chapter][section].strip()

    return structure

def updatetrainset(A, B,trainset,ca):
    A_structure = extract_sections(A)[0]
    print("A_structure",A_structure)
    B_structure = extract_sections(B)[0]
    print("B_structure", B_structure)
    for chapter in B_structure:
        if chapter not in A_structure:
            current_content=B_structure[chapter][0].split(":")[1].replace("\"","")
            current_content_segs = segment_text(current_content,2)
            label=chapter+","+B_structure[chapter][0].split(":")[0]
            for i in current_content_segs:
                ca.loc[len(ca)] = {'current_content': i, 'label1': label, 'label2': label}
                trainset.append(dspy.Example(
                    current_content=i,
                    label=label+",和"+label
                ).with_inputs('part_chapters', 'current_content'))

        else:
            for section in B_structure[chapter]:
                if section not in A_structure[chapter]:
                    print("not in a's section",section)
                    for asection in A_structure[chapter]:
                        if asection.split(":")[0]==section.split(":")[0]:
                            print("the same section is",asection)
                            current_content = difference(asection.split(":")[1].replace("\"",""),section.split(":")[1].replace("\"", ""))
                            current_content_segs=segment_text(current_content,2)
                            print("current_content_segs",current_content_segs)
                            for i in current_content_segs:
                                label = chapter + "," + section.split(":")[0]
                                ca.loc[len(ca)] = {'current_content': i, 'label1': label, 'label2': label}
                                trainset.append(dspy.Example(
                                    current_content=i,
                                    label=label+",和"+label).with_inputs('part_chapters', 'current_content'))
    return trainset,ca


def extract_content_and_label(text):
    pattern = re.compile(r"(.*?)\((Chapter[^)]+)\)")
    matches = pattern.findall(text)
    return [(content.strip(), label.strip()) for content, label in matches]

def updatetrainset2(text1, text2,trainset):
    content_label_pairs1 = extract_content_and_label(text1)
    content_label_pairs2 = extract_content_and_label(text2)

    differences = []
    for (content1, label1), (content2, label2) in zip(content_label_pairs1, content_label_pairs2):
        if label1 != label2:
            differences.append((content1, label1))
    for content, label in differences:
        trainset.append(dspy.Example(
            current_content=content,
            label=label + ",和" + label).with_inputs('part_chapters', 'current_content'))
    return trainset


    return differences

def parse_string(input_string):
    lines = input_string.split('\n')
    stack = []
    root = None
    current_node = None

    for line in lines:
        line = line.strip()
        if line.startswith('Part'):
            root = Node('Part Chapters')
            stack.append(root)
            current_node = root
        elif line.startswith('Chapter'):
            chapter_name = line
            chapter_node = Node(chapter_name)
            current_node.children.append(chapter_node)
            stack.append(chapter_node)
            current_node = chapter_node
        elif line.startswith('Section'):
            section_parts = line.split(':')
            section_name = section_parts[0]
            section_content = section_parts[1].strip().replace("\"", "") if len(section_parts) > 1 else ""
            section_node = Node(section_name)
            section_node.add_content(section_content)
            current_node.children.append(section_node)

    return root

def genprompt(tag,file):
    dfprompt=pd.read_csv(file)
    tagprompt=""
    for i in range(1,len(dfprompt.columns)):
        try:
            tagprompt+=dfprompt.columns[i]
            tagprompt+=" is "
            tagprompt+=dfprompt[dfprompt['tag'] == tag][dfprompt.columns[i]].astype(str).values[0]
            tagprompt+="; "
        except:
            tagprompt+=""
    str(tagprompt)
    return str(tagprompt)

def llmwriting(prompt):
    client = OpenAI(api_key="sk-proj-v61lipjKs2xKR7PXE2KdT3BlbkFJHyVAFIm8ilz01VjmFk9Z",)
    completion = client.chat.completions.create(
        model="gpt-4o-2024-05-13",
        messages=[{"role": "system", "content": prompt}]
    )
    return completion.choices[0].message.content

def format_result_list(result_list):
    formatted_output = ""
    for item in result_list:
        chapter = item.get('chapter', '')
        section = item.get('section', '')
        written = item.get('written', '')

        if chapter:
            formatted_output += f"**{chapter}**\n\n"
        if section:
            formatted_output += f"**Section: {section}**\n\n"
        formatted_output += f"{written}\n\n"

    return formatted_output.strip()


def merge_sections(node):
    if not node.children:
        return
    # Create a dictionary to store the merged nodes
    merged_children = {}

    for chapter in node.children:
        for section in chapter.children:
            # Step 1: Clean section.name
            clean_name = section.name.replace("[Not Sure]", "")

            # Step 2: Merge sections with the same cleaned name
            if clean_name not in merged_children:
                # If the name doesn't exist in the dictionary, add it and copy content
                merged_children[clean_name] = Node(clean_name)
                merged_children[clean_name].add_content(section.content)
            else:
                # If the name already exists, merge the content
                merged_children[clean_name].add_content(section.content)
        chapter.children=list(merged_children.values())
    print("merged children values",merged_children.values())
    return node  # Return the updated root node


def generate_writting(root, resultlist,file):
    new_root = copy.deepcopy(root)
    new_root = merge_sections(new_root)
    allcontent=tree_to_string(new_root)
    print("allcontent",allcontent)
    summary=llmwriting("What's the scenario of the following conversation? Summarize in one sentence:"+allcontent+"用中文")
    print("summary",summary)
    for item in resultlist:
        for chapter in new_root.children:
            print("aaaaa啊啊啊chapter啊啊啊",chapter)
            for section in chapter.children:
                print("bbbsection",section)
                chapter_name = chapter.name.split("Chapter")[1]
                section_name = section.name.split("Section")[1].replace("[Not Sure]","")
                if chapter_name==item['chapter'] and section_name==item['section']:
                    print("section content",section.content)
                    writing_style = item['writingStyle']
                    prompt = genprompt(writing_style, file)
                    print("writingstyle is:",writing_style, "prompt is:", prompt)
                    item['written']= llmwriting("input is "+section.content+prompt+"Select the part in the input that relevant to "+section_name+", and finish this part's writing of the whole conversation, which is about:" +summary +"Your output will directly be used as part of the report, just output it as a paragraph.")
                    print("this time: "+"input is "+section.content+prompt)
                    break
    print("after writing allocation",resultlist)
    output = ""
    for item in resultlist:
        output += f"**{item['chapter']}**\n"
        output += f"*{item['section']}*\n"
        try:
            output += item['written'] + "\n"
        except:
            output += "No content"
        output += "\n"
    if len(str(resultlist))>1000:
        return output
    else:
        return llmwriting("Complete the blank sections and output the whole formatted passage. Make sure each section has the contents align with the section names. Don't output the json list. Your output will be directly used as the final report."+str(resultlist))

   # return format_result_list(resultlist)

def searchnewchapter(new_root, resultlist):
    newsectionlist = []
    for chapter in new_root.children:
        chapter_name = chapter.name.split("Chapter")[1]
        for section in chapter.children:
            section_name = section.name.split("Section")[1]
            writing_style = None
            for result in resultlist:
                if result['chapter'] == chapter_name and result['section'] == section_name:
                    writing_style = result['writingStyle']
                    break
            if writing_style is None:
                newsectionlist.append(f"A new section with chapter name:{chapter_name}, and section name:{section_name}, and it should be written with xxx.")
    return newsectionlist


def updateresultlist(resultList, updatedchapterliststr):
    pattern = re.compile(r"chapter name:(.*?), and section name:(.*?), and it should be written with (.*?)(?:[,.?]|$)")
    matches = pattern.findall(updatedchapterliststr)
    for match in matches:
        chapter_name = match[0].strip()
        section_name = match[1].strip()
        writing_style = match[2].strip()
        exists = any(entry['chapter'] == chapter_name and entry['section'] == section_name for entry in resultList)

        if not exists:
            new_entry = {
                'chapter': chapter_name,
                'section': section_name,
                'writingStyle': writing_style
            }
            resultList.append(new_entry)

    return resultList


def updatetewritingstyle(A, B, ws, resultlist):
    ws['difference'] = ws['difference'].astype(str)
    A_structure = extract_sections(A)
    print("A_structure",A_structure)
    B_structure = extract_sections(B)
    print("B_structure", B_structure)

    def find_writing_style(chaptername, sectionname, resultlist):
        for item in resultlist:
            if item['chapter'] == chaptername and item['section'] == sectionname:

                return item['writingStyle']
        return None

    for chapter in B_structure:
        chaptername = chapter.split(":")[0].replace("**", "")
        print("chaptername",chaptername)
        if chapter in A_structure:
            for section in B_structure[chapter]:
                sectionname = section.split(":")[0].replace("*", "")
                print("sectionname",sectionname)
                if section in A_structure[chapter]:
                    if A_structure[chapter][section].strip() != B_structure[chapter][section].strip():
                        print("difference name",A_structure[chapter][section].strip())
                        current_writingstyle = find_writing_style(chaptername, sectionname, resultlist)
                        ws.loc[ws['tag'] == current_writingstyle, 'difference'] = str(llmwriting("The previous version is: "+A_structure[chapter][section].strip()+". The latter version is: " +B_structure[chapter][section].strip()+". Compare the writing styles of these versions. Summarize the changes made in the latter version and suggest how to improve the writing style to be more like the latter version. Avoid including detailed information; focus on summarizing the overall writing styles.")) ##不要写具体东西

    return ws

def toJson(example_instance):
    # Convert _store to a JSON-serializable dictionary
    serializable_data = {}
    for key, value in example_instance._store.items():
        if isinstance(value, dspy.Example):
            serializable_data[key] = toJson(value)  # Recursively convert nested Example instances
        else:
            serializable_data[key] = value

    # Split "label" into "label1" and "label2" if present
    if "label" in serializable_data:
        serializable_data["label"]=serializable_data["label"].replace("\\","").replace("\"","")
        labels = serializable_data["label"].split(",和")
        serializable_data["label1"] = labels[0]
        if len(labels) > 1:
            serializable_data["label2"] = labels[1]
        del serializable_data["label"]

    return json.dumps(serializable_data, default=str, ensure_ascii=False)